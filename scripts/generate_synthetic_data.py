"""Generate deterministic synthetic NovaPromo data (no real company data is used).

Outputs
  data/suppliers/supplier_1.xml   EuroGift Trading GmbH (DE) - XML, comma decimals, German colours
  data/suppliers/supplier_2.json  PromoTex s.r.o. (CZ)       - JSON, textile, cents, size variants
  data/suppliers/supplier_3.xlsx  EcoLine Supplies s.r.o.(CZ)- XLSX, Czech headers, CZK prices
  data/products/sku_crosswalk.csv ERP mapping supplier code -> NovaPromo SKU
  data/products/warehouse_stock.csv own warehouse stock (Bratislava / Zilina)
  data/products/products_normalized.csv  the normalised catalogue (output of the ingestion pipeline)

Usage: python scripts/generate_synthetic_data.py
"""

from __future__ import annotations

import csv
import json
import random
import sys
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from packages.core.text import COLOR_NAMES, MATERIAL_NAMES  # noqa: E402

SUPPLIERS_DIR = ROOT / "data" / "suppliers"
PRODUCTS_DIR = ROOT / "data" / "products"
CZK_PER_EUR = 25.0
rng = random.Random(42)

GERMAN_COLOURS = {
    "black": "Schwarz",
    "white": "Weiss",
    "blue": "Blau",
    "red": "Rot",
    "green": "Grün",
    "grey": "Grau",
    "silver": "Silber",
    "navy": "Marineblau",
    "yellow": "Gelb",
    "orange": "Orange",
}
PROMOTEX_CODES = {
    "black": "BK",
    "white": "WH",
    "navy": "NV",
    "blue": "RB",
    "red": "RD",
    "grey": "GR",
    "green": "GN",
    "yellow": "YL",
    "natural": "NT",
    "orange": "OR",
}

BRANDING = {
    "laser": ("laser engraving", "laserové gravírovanie", "laserové gravírování"),
    "screen": ("screen printing", "sieťotlač", "sítotisk"),
    "pad": ("pad printing", "tampónová tlač", "tampontisk"),
    "embroidery": ("embroidery", "výšivka", "výšivka"),
    "uv": ("full-colour UV print", "plnofarebná UV tlač", "plnobarevný UV tisk"),
    "transfer": ("transfer print", "transferová tlač", "transferový tisk"),
}
USE_CASES = {
    "corporate": (
        "corporate gifts and client welcome packs",
        "firemné darčeky a uvítacie balíčky pre klientov",
        "firemní dárky a uvítací balíčky pro klienty",
    ),
    "event": (
        "conferences, trade fairs and events",
        "konferencie, veľtrhy a podujatia",
        "konference, veletrhy a akce",
    ),
    "outdoor": (
        "outdoor activities and travel",
        "outdoorové aktivity a cestovanie",
        "outdoorové aktivity a cestování",
    ),
    "office": (
        "the office and employee onboarding",
        "kanceláriu a onboarding zamestnancov",
        "kancelář a onboarding zaměstnanců",
    ),
    "eco": (
        "sustainability-focused campaigns",
        "kampane zamerané na udržateľnosť",
        "kampaně zaměřené na udržitelnost",
    ),
    "tech": (
        "technology companies and IT events",
        "technologické firmy a IT podujatia",
        "technologické firmy a IT akce",
    ),
}


@dataclass
class ProductType:
    key: str
    supplier: int
    names: tuple[str, str, str]  # en, sk, cs
    category: str
    subcategory: str
    materials: list[str]
    price: tuple[float, float]
    colors: list[str]
    models: list[str]
    features: tuple[str, str, str]
    uses: list[str]
    branding: list[str]
    tags: list[str]
    capacities: list[int] = field(default_factory=list)
    brands: list[str] = field(default_factory=lambda: ["NovaLine"])
    origin: list[str] = field(default_factory=lambda: ["CN"])
    sizes: list[str] = field(default_factory=list)
    lead: tuple[int, int] = (5, 14)


TYPES: list[ProductType] = [
    # ---------------- Supplier 1: EuroGift Trading (XML) -------------------------------------
    ProductType(
        "thermal_bottle",
        1,
        ("Thermal bottle", "Termofľaša", "Termoláhev"),
        "drinkware",
        "bottles",
        ["stainless steel"],
        (8.5, 22.0),
        ["black", "white", "blue", "silver", "red", "green", "navy"],
        ["Arctic", "Nordic", "Summit", "Glacier", "Tatra"],
        (
            "Double-wall vacuum insulation keeps drinks hot for 12 hours and cold for 24 hours",
            "Dvojstenná vákuová izolácia udrží nápoje teplé 12 hodín a studené 24 hodín",
            "Dvoustěnná vakuová izolace udrží nápoje teplé 12 hodin a studené 24 hodin",
        ),
        ["corporate", "outdoor"],
        ["laser", "pad"],
        ["insulated", "vacuum", "bpa-free", "leak-proof"],
        capacities=[500, 750, 1000],
        brands=["Thermaline", "NovaLine"],
        origin=["CN", "TW"],
    ),
    ProductType(
        "sports_bottle",
        1,
        ("Sports bottle", "Športová fľaša", "Sportovní láhev"),
        "drinkware",
        "bottles",
        ["aluminium"],
        (3.4, 8.9),
        ["black", "white", "blue", "red", "silver", "green", "orange"],
        ["Sprint", "Pace", "Velo"],
        (
            "Lightweight single-wall aluminium bottle with a carabiner lid",
            "Ľahká jednostenná hliníková fľaša s viečkom s karabínou",
            "Lehká jednostěnná hliníková láhev s víčkem s karabinou",
        ),
        ["outdoor", "event"],
        ["pad", "laser"],
        ["lightweight", "sport", "bpa-free"],
        capacities=[400, 600, 750],
        brands=["ActiveGo"],
    ),
    ProductType(
        "glass_bottle",
        1,
        (
            "Glass bottle with bamboo lid",
            "Sklenená fľaša s bambusovým viečkom",
            "Skleněná láhev s bambusovým víčkem",
        ),
        "drinkware",
        "bottles",
        ["glass"],
        (4.9, 11.5),
        ["natural", "black", "blue", "green"],
        ["Clarity", "Aqua"],
        (
            "Borosilicate glass bottle with a bamboo lid and a neoprene sleeve",
            "Fľaša z borosilikátového skla s bambusovým viečkom a neoprénovým obalom",
            "Láhev z borosilikátového skla s bambusovým víčkem a neoprenovým obalem",
        ),
        ["eco", "office"],
        ["laser", "screen"],
        ["eco", "plastic-free", "bamboo"],
        capacities=[500, 600],
        brands=["GreenWave"],
    ),
    ProductType(
        "ceramic_mug",
        1,
        ("Ceramic mug", "Keramický hrnček", "Keramický hrnek"),
        "drinkware",
        "mugs",
        ["ceramic"],
        (2.4, 6.9),
        ["white", "black", "blue", "red", "green", "yellow", "grey"],
        ["Classic", "Barista", "Studio", "Metro"],
        (
            "Dishwasher-safe ceramic mug with a glossy glaze",
            "Keramický hrnček vhodný do umývačky s lesklou glazúrou",
            "Keramický hrnek vhodný do myčky s lesklou glazurou",
        ),
        ["office", "corporate"],
        ["screen", "transfer"],
        ["dishwasher-safe", "office"],
        capacities=[300, 350, 400],
        brands=["NovaLine"],
        origin=["PT", "CN"],
    ),
    ProductType(
        "tumbler",
        1,
        ("Travel tumbler", "Cestovný pohár", "Cestovní hrnek"),
        "drinkware",
        "tumblers",
        ["stainless steel"],
        (6.9, 15.9),
        ["black", "white", "silver", "navy", "red"],
        ["Commuter", "Metro"],
        (
            "Leak-resistant travel tumbler that fits most car cup holders",
            "Cestovný pohár odolný proti vytečeniu, ktorý sa zmestí do väčšiny držiakov v aute",
            "Cestovní hrnek odolný proti vytečení, který se vejde do většiny držáků v autě",
        ),
        ["corporate", "outdoor"],
        ["laser"],
        ["insulated", "coffee", "travel"],
        capacities=[350, 450],
        brands=["Thermaline"],
    ),
    ProductType(
        "umbrella",
        1,
        ("Storm umbrella", "Dáždnik odolný voči búrke", "Deštník odolný bouři"),
        "umbrellas",
        "umbrellas",
        ["polyester"],
        (6.5, 24.9),
        ["black", "navy", "red", "blue", "grey", "yellow"],
        ["Storm", "Breeze", "City"],
        (
            "Windproof fibreglass frame and automatic opening",
            "Vetruodolná konštrukcia zo sklolaminátu a automatické otváranie",
            "Větruodolná konstrukce ze skelného vlákna a automatické otevírání",
        ),
        ["outdoor", "corporate"],
        ["screen", "transfer"],
        ["windproof", "automatic"],
        brands=["RainPro"],
    ),
    ProductType(
        "backpack",
        1,
        ("Laptop backpack", "Batoh na notebook", "Batoh na notebook"),
        "bags",
        "backpacks",
        ["polyester", "recycled pet"],
        (14.0, 45.0),
        ["black", "grey", "navy", "blue"],
        ["Metro", "Voyager", "Pixel", "Campus"],
        (
            "Padded compartment for 15.6-inch laptops and an anti-theft back pocket",
            "Polstrované priehradky pre notebooky do 15,6 palca a zadné vrecko proti krádeži",
            "Polstrovaná přihrádka pro notebooky do 15,6 palce a zadní kapsa proti krádeži",
        ),
        ["office", "tech", "outdoor"],
        ["embroidery", "transfer"],
        ["laptop", "travel", "anti-theft"],
        brands=["Urbanpack"],
        lead=(7, 21),
    ),
    ProductType(
        "conference_bag",
        1,
        ("Conference bag", "Konferenčná taška", "Konferenční taška"),
        "bags",
        "conference bags",
        ["polyester", "recycled pet", "jute"],
        (2.1, 8.9),
        ["black", "navy", "natural", "grey", "blue"],
        ["Delegate", "Expo", "Summit"],
        (
            "Document bag with an adjustable shoulder strap and a front pocket for badges",
            "Taška na dokumenty s nastaviteľným popruhom a predným vreckom na menovky",
            "Taška na dokumenty s nastavitelným popruhem a přední kapsou na jmenovky",
        ),
        ["event"],
        ["screen", "transfer"],
        ["conference", "event", "documents"],
        brands=["NovaLine"],
    ),
    ProductType(
        "metal_pen",
        1,
        ("Metal ballpoint pen", "Kovové guľôčkové pero", "Kovové kuličkové pero"),
        "writing",
        "pens",
        ["aluminium"],
        (0.6, 3.5),
        ["black", "silver", "blue", "red", "green", "navy"],
        ["Signature", "Slim", "Orion"],
        (
            "Twist mechanism, blue ink refill and a soft-touch finish",
            "Otočný mechanizmus, modrá náplň a povrch soft-touch",
            "Otočný mechanismus, modrá náplň a povrch soft-touch",
        ),
        ["office", "event"],
        ["laser"],
        ["office", "writing"],
        brands=["Scriba"],
    ),
    ProductType(
        "power_bank",
        1,
        ("Power bank", "Powerbanka", "Powerbanka"),
        "technology",
        "power banks",
        ["abs", "aluminium"],
        (8.9, 29.9),
        ["black", "white", "silver", "blue"],
        ["Volt", "Pulse", "Charge"],
        (
            "Lithium-polymer battery with USB-C input/output and LED charge indicator",
            "Lítium-polymérová batéria s USB-C vstupom/výstupom a LED indikátorom nabitia",
            "Lithium-polymerová baterie s USB-C vstupem/výstupem a LED indikátorem nabití",
        ),
        ["tech", "corporate"],
        ["uv", "laser"],
        ["usb-c", "tech", "charging"],
        brands=["VoltEdge"],
    ),
    ProductType(
        "wireless_charger",
        1,
        ("Wireless charger", "Bezdrôtová nabíjačka", "Bezdrátová nabíječka"),
        "technology",
        "chargers",
        ["bamboo", "abs"],
        (8.5, 21.9),
        ["natural", "black", "white"],
        ["Qi Pad", "Desk Pad"],
        (
            "10 W Qi-certified wireless charging pad for smartphones",
            "10 W bezdrôtová nabíjacia podložka s certifikáciou Qi pre smartfóny",
            "10W bezdrátová nabíjecí podložka s certifikací Qi pro chytré telefony",
        ),
        ["tech", "office"],
        ["laser", "uv"],
        ["wireless", "qi", "tech"],
        brands=["VoltEdge"],
    ),
    ProductType(
        "speaker",
        1,
        ("Bluetooth speaker", "Bluetooth reproduktor", "Bluetooth reproduktor"),
        "technology",
        "speakers",
        ["abs"],
        (11.9, 34.9),
        ["black", "blue", "red", "grey"],
        ["Beat", "Echo"],
        (
            "Bluetooth 5.3 speaker with 8 hours of playback and a built-in microphone",
            "Bluetooth 5.3 reproduktor s 8 hodinami prehrávania a zabudovaným mikrofónom",
            "Bluetooth 5.3 reproduktor s 8 hodinami přehrávání a vestavěným mikrofonem",
        ),
        ["tech", "outdoor"],
        ["uv", "pad"],
        ["bluetooth", "wireless", "tech"],
        brands=["VoltEdge"],
    ),
    # ---------------- Supplier 2: PromoTex (JSON) --------------------------------------------
    ProductType(
        "tshirt",
        2,
        ("T-shirt", "Tričko", "Tričko"),
        "apparel",
        "t-shirts",
        ["cotton", "organic cotton"],
        (2.9, 9.5),
        ["black", "white", "navy", "blue", "red", "grey", "green", "yellow"],
        ["Basic 150", "Premium 180", "Heavy 200", "Organic 160", "Fit"],
        (
            "Pre-shrunk jersey with a taped neck and side seams",
            "Vopred zrazený úplet s páskovaným výstrihom a bočnými švami",
            "Předem sražený úplet s podlepeným výstřihem a bočními švy",
        ),
        ["event", "corporate"],
        ["screen", "embroidery", "transfer"],
        ["textile", "unisex"],
        sizes=["XS", "S", "M", "L", "XL", "XXL"],
        brands=["TexLine", "Organico"],
        origin=["BD", "TR", "PT"],
    ),
    ProductType(
        "polo",
        2,
        ("Polo shirt", "Polokošeľa", "Polokošile"),
        "apparel",
        "polo shirts",
        ["cotton"],
        (6.9, 15.9),
        ["black", "white", "navy", "blue", "red", "grey"],
        ["Pique", "Club", "Office"],
        (
            "Piqué knit with a three-button placket and ribbed collar",
            "Piké úplet s légou na tri gombíky a rebrovaným golierom",
            "Piké úplet s légou na tři knoflíky a žebrovaným límcem",
        ),
        ["corporate", "event"],
        ["embroidery"],
        ["textile", "workwear"],
        sizes=["S", "M", "L", "XL", "XXL", "3XL"],
        brands=["TexLine"],
        origin=["BD", "TR"],
    ),
    ProductType(
        "hoodie",
        2,
        ("Hoodie", "Mikina s kapucňou", "Mikina s kapucí"),
        "apparel",
        "hoodies",
        ["cotton", "organic cotton", "polyester"],
        (13.9, 32.0),
        ["black", "grey", "navy", "white", "green"],
        ["Campus", "Urban", "Cozy"],
        (
            "Brushed fleece inside, kangaroo pocket and a double-layer hood",
            "Vnútri počesaný flís, klokanie vrecko a dvojvrstvová kapucňa",
            "Uvnitř počesaný fleece, klokaní kapsa a dvouvrstvá kapuce",
        ),
        ["tech", "corporate"],
        ["embroidery", "screen"],
        ["textile", "winter"],
        sizes=["S", "M", "L", "XL", "XXL"],
        brands=["TexLine", "Organico"],
        origin=["BD", "PK"],
    ),
    ProductType(
        "cap",
        2,
        ("Baseball cap", "Šiltovka", "Kšiltovka"),
        "apparel",
        "caps",
        ["cotton"],
        (1.9, 6.4),
        ["black", "white", "navy", "red", "blue", "grey", "yellow"],
        ["Classic 5", "Sport 6"],
        (
            "Brushed cotton twill with a metal buckle closure",
            "Počesaný bavlnený kepr s kovovou prackou",
            "Počesaný bavlněný kepr s kovovou přezkou",
        ),
        ["event", "outdoor"],
        ["embroidery"],
        ["textile", "summer"],
        brands=["TexLine"],
        origin=["CN"],
    ),
    ProductType(
        "cotton_tote",
        2,
        ("Cotton tote bag", "Bavlnená taška", "Bavlněná taška"),
        "bags",
        "tote bags",
        ["cotton", "organic cotton"],
        (1.2, 4.6),
        ["natural", "black", "navy", "red", "green"],
        ["Market", "Shopper", "Eco Tote"],
        (
            "140 g/m² cotton with long handles, reusable and washable",
            "Bavlna 140 g/m² s dlhými ušami, opakovane použiteľná a prateľná",
            "Bavlna 140 g/m² s dlouhými uchy, opakovaně použitelná a pratelná",
        ),
        ["eco", "event"],
        ["screen", "transfer"],
        ["eco", "reusable", "textile"],
        brands=["Organico"],
        origin=["IN"],
    ),
    ProductType(
        "lanyard",
        2,
        ("Lanyard", "Šnúrka na krk", "Šňůrka na krk"),
        "lifestyle",
        "lanyards",
        ["polyester", "recycled pet"],
        (0.35, 1.6),
        ["black", "blue", "red", "green", "white", "navy"],
        ["Badge 20", "Safety 15"],
        (
            "Safety breakaway clip and a metal hook for badges",
            "Bezpečnostná rozpojovacia spona a kovový karabínok na menovky",
            "Bezpečnostní rozpojovací spona a kovová karabinka na jmenovky",
        ),
        ["event"],
        ["transfer"],
        ["conference", "event", "badge"],
        brands=["TexLine"],
        origin=["CN"],
    ),
    # ---------------- Supplier 3: EcoLine (XLSX) ---------------------------------------------
    ProductType(
        "bamboo_notebook",
        3,
        ("Bamboo cover notebook A5", "Zápisník s bambusovým obalom A5", "Zápisník s bambusovým obalem A5"),
        "office",
        "notebooks",
        ["bamboo"],
        (3.2, 9.4),
        ["natural"],
        ["Terra", "Forest", "Leaf"],
        (
            "80 sheets of lined recycled paper and an elastic closure band",
            "80 listov linajkového recyklovaného papiera a elastická gumička",
            "80 listů linkovaného recyklovaného papíru a elastická gumička",
        ),
        ["eco", "office", "event"],
        ["laser"],
        ["eco", "bamboo", "recycled"],
        brands=["EcoLine"],
        origin=["CN"],
    ),
    ProductType(
        "recycled_notebook",
        3,
        (
            "Recycled paper notebook A5",
            "Zápisník z recyklovaného papiera A5",
            "Zápisník z recyklovaného papíru A5",
        ),
        "office",
        "notebooks",
        ["recycled paper"],
        (1.4, 5.2),
        ["natural", "black", "blue", "green", "red"],
        ["Kraft", "Loop"],
        (
            "Kraft cardboard cover, FSC-certified paper and a matching pen loop",
            "Obal z kraftového kartónu, papier s certifikátom FSC a pútko na pero",
            "Obal z kraftového kartonu, papír s certifikací FSC a poutko na pero",
        ),
        ["eco", "event", "office"],
        ["screen", "pad"],
        ["eco", "recycled", "fsc"],
        brands=["EcoLine"],
        origin=["CZ", "PL"],
    ),
    ProductType(
        "cork_notebook",
        3,
        ("Cork notebook A6", "Korkový zápisník A6", "Korkový zápisník A6"),
        "office",
        "notebooks",
        ["cork"],
        (2.8, 6.9),
        ["natural"],
        ["Porto", "Lisboa"],
        (
            "Natural cork cover with 70 plain recycled pages",
            "Obal z prírodného korku so 70 čistými recyklovanými stranami",
            "Obal z přírodního korku se 70 čistými recyklovanými stranami",
        ),
        ["eco", "event"],
        ["laser", "screen"],
        ["eco", "cork", "natural"],
        brands=["EcoLine"],
        origin=["PT"],
    ),
    ProductType(
        "wheat_mug",
        3,
        ("Wheat straw mug", "Hrnček z pšeničnej slamy", "Hrnek z pšeničné slámy"),
        "drinkware",
        "mugs",
        ["wheat straw"],
        (2.1, 4.9),
        ["natural", "green", "blue", "grey"],
        ["Harvest", "Field"],
        (
            "Made from wheat-straw fibre composite, reusable and lightweight",
            "Vyrobený z kompozitu z vlákien pšeničnej slamy, opakovane použiteľný a ľahký",
            "Vyrobený z kompozitu z vláken pšeničné slámy, opakovaně použitelný a lehký",
        ),
        ["eco", "office"],
        ["pad"],
        ["eco", "biodegradable", "coffee"],
        capacities=[300, 350],
        brands=["EcoLine"],
        origin=["CN"],
    ),
    ProductType(
        "rpet_bottle",
        3,
        ("rPET drinking bottle", "Fľaša z rPET", "Láhev z rPET"),
        "drinkware",
        "bottles",
        ["recycled pet"],
        (2.4, 5.9),
        ["blue", "green", "natural", "black", "grey"],
        ["Loop", "Refill"],
        (
            "Made from 100 % recycled PET, BPA-free with a leak-proof screw cap",
            "Vyrobená zo 100 % recyklovaného PET, bez BPA so skrutkovacím uzáverom proti vytečeniu",
            "Vyrobená ze 100 % recyklovaného PET, bez BPA se šroubovacím uzávěrem proti vytečení",
        ),
        ["eco", "event", "outdoor"],
        ["pad", "screen"],
        ["eco", "recycled", "bpa-free"],
        capacities=[500, 600, 750],
        brands=["EcoLine"],
        origin=["CZ"],
    ),
    ProductType(
        "plantable_pencil",
        3,
        ("Plantable pencil", "Rastúca ceruzka", "Rostoucí tužka"),
        "writing",
        "pencils",
        ["wood"],
        (0.7, 1.9),
        ["natural"],
        ["Sprout", "Seed"],
        (
            "Graphite pencil with a seed capsule - plant it when it gets short",
            "Grafitová ceruzka s kapsulou semienok - keď sa skráti, zasaďte ju",
            "Grafitová tužka s kapslí semínek - když se zkrátí, zasaďte ji",
        ),
        ["eco", "event"],
        ["pad", "laser"],
        ["eco", "wood", "giveaway"],
        brands=["EcoLine"],
        origin=["DK", "CZ"],
    ),
    ProductType(
        "jute_bag",
        3,
        ("Jute shopping bag", "Jutová nákupná taška", "Jutová nákupní taška"),
        "bags",
        "tote bags",
        ["jute"],
        (2.6, 7.2),
        ["natural"],
        ["Market", "Harvest"],
        (
            "Laminated jute with cotton handles, durable and reusable",
            "Laminovaná juta s bavlnenými ušami, odolná a opakovane použiteľná",
            "Laminovaná juta s bavlněnými uchy, odolná a opakovaně použitelná",
        ),
        ["eco", "event"],
        ["screen"],
        ["eco", "reusable", "natural"],
        brands=["EcoLine"],
        origin=["IN"],
    ),
    ProductType(
        "bamboo_charger",
        3,
        ("Bamboo wireless charger", "Bambusová bezdrôtová nabíjačka", "Bambusová bezdrátová nabíječka"),
        "technology",
        "chargers",
        ["bamboo"],
        (9.9, 19.9),
        ["natural"],
        ["Bamboo Qi", "Zen"],
        (
            "15 W fast wireless charging in a sustainably sourced bamboo housing",
            "15 W rýchle bezdrôtové nabíjanie v bambusovom kryte z udržateľných zdrojov",
            "15W rychlé bezdrátové nabíjení v bambusovém krytu z udržitelných zdrojů",
        ),
        ["eco", "tech", "office"],
        ["laser"],
        ["eco", "wireless", "tech", "bamboo"],
        brands=["EcoLine"],
        origin=["CN"],
    ),
    ProductType(
        "solar_powerbank",
        3,
        ("Solar power bank", "Solárna powerbanka", "Solární powerbanka"),
        "technology",
        "power banks",
        ["recycled pet", "abs"],
        (14.9, 29.9),
        ["black", "green", "grey"],
        ["Sunny", "Helio"],
        (
            "Solar panel top-up charging and a housing made from recycled plastic",
            "Dobíjanie solárnym panelom a kryt z recyklovaného plastu",
            "Dobíjení solárním panelem a kryt z recyklovaného plastu",
        ),
        ["eco", "tech", "outdoor"],
        ["uv", "pad"],
        ["eco", "solar", "tech", "usb-c"],
        brands=["EcoLine"],
        origin=["CN"],
    ),
    ProductType(
        "cork_coasters",
        3,
        ("Cork coaster set", "Sada korkových podložiek", "Sada korkových podtácků"),
        "lifestyle",
        "home",
        ["cork"],
        (2.9, 7.9),
        ["natural"],
        ["Duo", "Quattro"],
        (
            "Set of four round natural-cork coasters in a kraft gift box",
            "Sada štyroch okrúhlych podložiek z prírodného korku v darčekovej krabičke",
            "Sada čtyř kulatých podtácků z přírodního korku v dárkové krabičce",
        ),
        ["eco", "corporate"],
        ["laser"],
        ["eco", "cork", "gift"],
        brands=["EcoLine"],
        origin=["PT"],
    ),
]

DESC_TEMPLATES = {
    "en": "{name} made of {material}. {features}.{capacity} Ideal for {use}. Branding options: {branding}.",
    "sk": "{name} z materiálu {material}. {features}.{capacity} Ideálne pre {use}. Možnosti potlače: {branding}.",
    "cs": "{name} z materiálu {material}. {features}.{capacity} Ideální pro {use}. Možnosti potisku: {branding}.",
}
CAPACITY_TEMPLATES = {"en": " Capacity {ml} ml.", "sk": " Objem {ml} ml.", "cs": " Objem {ml} ml."}
LANG_IDX = {"en": 0, "sk": 1, "cs": 2}


def color_name(color: str, lang: str) -> str:
    if lang == "en":
        return color
    return COLOR_NAMES[color][0 if lang == "sk" else 1]


def material_name(material: str, lang: str) -> str:
    if lang == "en":
        return material
    return MATERIAL_NAMES[material][0 if lang == "sk" else 1]


def product_name(t: ProductType, model: str, color: str, cap: int | None, lang: str) -> str:
    base = f"{t.names[LANG_IDX[lang]]} {model}"
    if cap:
        base += f" {cap} ml"
    return f"{base}, {color_name(color, lang)}"


def description(t: ProductType, name: str, material: str, cap: int | None, lang: str) -> str:
    i = LANG_IDX[lang]
    use = " / ".join(USE_CASES[u][i] for u in t.uses[:2])
    branding = ", ".join(BRANDING[b][i] for b in t.branding)
    capacity = CAPACITY_TEMPLATES[lang].format(ml=cap) if cap else ""
    return DESC_TEMPLATES[lang].format(
        name=name.split(",")[0],
        material=material_name(material, lang),
        features=t.features[i],
        capacity=capacity,
        use=use,
        branding=branding,
    )


MODEL_POOL = [
    "Aurora",
    "Vega",
    "Nova",
    "Delta",
    "Atlas",
    "Luna",
    "Sierra",
    "Alpine",
    "Coast",
    "Danube",
    "Carpathia",
    "Quantum",
    "Zenith",
    "Horizon",
    "Fusion",
    "Polar",
    "Comet",
    "Lumen",
    "Rift",
    "Kite",
]
MODELS_PER_SUPPLIER = {1: 8, 2: 10, 3: 8}


def build_catalogue() -> list[dict]:
    items: list[dict] = []
    master_seq = {1: 30100, 2: 4400, 3: 100}
    for t in TYPES:
        caps = t.capacities or [None]
        extra = [m for m in rng.sample(MODEL_POOL, len(MODEL_POOL)) if m not in t.models]
        models = t.models + extra[: max(0, MODELS_PER_SUPPLIER[t.supplier] - len(t.models))]
        for model in models:
            for cap in caps:
                material = rng.choice(t.materials)
                master_seq[t.supplier] += 1 if t.supplier != 2 else 10
                master = master_seq[t.supplier]
                n_colors = min(len(t.colors), rng.randint(2, 6))
                colors = rng.sample(t.colors, n_colors)
                if t.key == "thermal_bottle" and model == "Arctic" and cap == 750:
                    master = 29904  # hero master code used in demos
                    colors = ["black"] + [c for c in colors if c != "black"][:4]
                    material = "stainless steel"
                base = round(rng.uniform(*t.price) * (1 + ((cap or 500) - 500) / 2500), 2)
                for vi, color in enumerate(colors, start=1):
                    price = max(0.3, round(base * rng.uniform(0.95, 1.05), 2))
                    if t.key == "thermal_bottle" and master == 29904 and color == "black":
                        price = 12.90
                    if t.supplier == 1:
                        sku = f"{master}.{vi:02d}"
                    elif t.supplier == 2:
                        sku = f"PT-{master}-{PROMOTEX_CODES[color]}"
                    else:
                        sku = f"ECO-{master:04d}-{vi}"
                    names = {lang: product_name(t, model, color, cap, lang) for lang in ("en", "sk", "cs")}
                    descs = {
                        lang: description(t, names[lang], material, cap, lang) for lang in ("en", "sk", "cs")
                    }
                    tags = sorted(set(t.tags + rng.sample(t.tags, min(1, len(t.tags)))))
                    item = {
                        "type": t.key,
                        "supplier": t.supplier,
                        "supplier_sku": sku,
                        "master": master,
                        "model": model,
                        "names": names,
                        "descs": descs,
                        "category": t.category,
                        "subcategory": t.subcategory,
                        "color": color,
                        "material": material,
                        "capacity_ml": cap,
                        "sizes": t.sizes,
                        "price_eur": price,
                        "brand": rng.choice(t.brands),
                        "origin": rng.choice(t.origin),
                        "lead": rng.randint(*t.lead),
                        "stock": rng.choice([0, 0, 40, 120, 250, 600, 1200, 3500, 8000]),
                        "tags": tags,
                    }
                    if item["master"] == 29904 and color == "black":
                        item["stock"] = 1200
                    items.append(item)
    return items


# ----------------------------------------------------------------------------------------------
def write_supplier_1(items: list[dict]) -> None:
    cat_paths = {
        "bottles": "Drinkware > Vacuum Flasks & Bottles",
        "mugs": "Drinkware > Mugs",
        "tumblers": "Drinkware > Travel Mugs",
        "umbrellas": "Rain > Umbrellas",
        "backpacks": "Bags > Rucksacks",
        "conference bags": "Bags > Document Bags",
        "pens": "Writing > Ballpens",
        "power banks": "Electronics > Powerbanks",
        "chargers": "Electronics > Chargers",
        "speakers": "Electronics > Audio",
    }
    root = ET.Element("catalog", supplier="EGT", currency="EUR", generated="2026-09-30T22:00:00Z")
    for it in items:
        el = ET.SubElement(root, "item", code=it["supplier_sku"])
        ET.SubElement(el, "master_code").text = str(it["master"])
        names = ET.SubElement(el, "names")
        for lang in ("en", "sk"):
            ET.SubElement(names, "name", lang=lang).text = it["names"][lang]
        descs = ET.SubElement(el, "descriptions")
        ET.SubElement(descs, "description", lang="en").text = it["descs"]["en"]
        if rng.random() < 0.55:  # ~45 % of items lack a Slovak description -> batch job fills them
            ET.SubElement(descs, "description", lang="sk").text = it["descs"]["sk"]
        ET.SubElement(el, "category", path=cat_paths[it["subcategory"]])
        attrs = ET.SubElement(el, "attributes")
        colour = (
            GERMAN_COLOURS.get(it["color"], it["color"].title())
            if rng.random() < 0.3
            else it["color"].title()
        )
        if it.get("_broken") == "colour":
            colour = ""
        ET.SubElement(attrs, "attr", name="colour").text = colour
        ET.SubElement(attrs, "attr", name="material").text = it["material"].title()
        if it["capacity_ml"]:
            vol = (
                f"{it['capacity_ml'] / 1000:.2f} L".replace(".", ",")
                if rng.random() < 0.3
                else f"{it['capacity_ml']} ml"
            )
            ET.SubElement(attrs, "attr", name="volume").text = vol
        pricing = ET.SubElement(el, "pricing", currency="EUR")
        price_txt = f"{it['price_eur']:.2f}".replace(".", ",")
        if it.get("_broken") == "price":
            price_txt = "N/A"
        ET.SubElement(pricing, "price", qty="1").text = price_txt
        if it.get("_broken") != "price":
            ET.SubElement(pricing, "price", qty="100").text = f"{it['price_eur'] * 0.92:.2f}".replace(
                ".", ","
            )
            ET.SubElement(pricing, "price", qty="500").text = f"{it['price_eur'] * 0.85:.2f}".replace(
                ".", ","
            )
        ET.SubElement(el, "stock", warehouse="DE-HAM").text = str(it["stock"])
        ET.SubElement(el, "brand").text = it["brand"]
        ET.SubElement(el, "origin").text = it["origin"]
        ET.SubElement(el, "leadtime", unit="days").text = str(it["lead"])
        ET.SubElement(el, "keywords").text = ";".join(it["tags"])
    ET.indent(root)
    ET.ElementTree(root).write(SUPPLIERS_DIR / "supplier_1.xml", encoding="utf-8", xml_declaration=True)


def write_supplier_2(items: list[dict]) -> None:
    groups = {
        "t-shirts": "T-Shirts",
        "polo shirts": "Polos",
        "hoodies": "Sweats & Hoodies",
        "caps": "Headwear",
        "tote bags": "Bags",
        "lanyards": "Accessories",
    }
    products = []
    for it in items:
        gsm = rng.choice([140, 150, 160, 180, 200])
        fabric = {
            "cotton": f"100% cotton, {gsm} g/m2",
            "organic cotton": f"100% organic cotton, {gsm} g/m2",
            "polyester": "100% polyester",
            "recycled pet": "100% recycled polyester (rPET)",
        }[it["material"]]
        sizes = it["sizes"] or ["ONE"]
        per = max(0, it["stock"] // len(sizes))
        body: dict[str, str] = {"en": it["descs"]["en"]}
        if rng.random() < 0.35:
            body["cs"] = it["descs"]["cs"]
        amount = int(round(it["price_eur"] * 100))
        if it.get("_broken") == "price":
            amount = -int(amount)
        products.append(
            {
                "id": it["supplier_sku"],
                "style": f"PT-{it['master']}",
                "title": {"en": it["names"]["en"], "cs": it["names"]["cs"], "sk": it["names"]["sk"]},
                "body": body,
                "group": groups[it["subcategory"]],
                "colour_code": PROMOTEX_CODES[it["color"]],
                "fabric": fabric,
                "sizes": [{"size": s, "stock": per} for s in sizes],
                "price": {"amount_cents": amount, "currency": "EUR"},
                "brand": it["brand"],
                "made_in": it["origin"],
                "delivery_days": it["lead"],
                "labels": it["tags"],
            }
        )
    feed = {
        "feed_version": "2.1",
        "supplier": "PromoTex s.r.o.",
        "exported_at": "2026-09-30T21:15:00+02:00",
        "colour_codes": {v: k for k, v in PROMOTEX_CODES.items()},
        "products": products,
    }
    (SUPPLIERS_DIR / "supplier_2.json").write_text(
        json.dumps(feed, ensure_ascii=False, indent=2), encoding="utf-8"
    )


def write_supplier_3(items: list[dict]) -> None:
    cats = {
        "notebooks": "Kancelář / Zápisníky",
        "mugs": "Nápoje / Hrnky",
        "bottles": "Nápoje / Láhve",
        "pencils": "Psací potřeby / Tužky",
        "tote bags": "Tašky / Nákupní tašky",
        "chargers": "Elektronika / Nabíječky",
        "power banks": "Elektronika / Powerbanky",
        "home": "Domov / Doplňky",
    }
    wb = Workbook()
    ws = wb.active
    ws.title = "Produkty"
    headers = [
        "Kód zboží",
        "Název CZ",
        "Název SK",
        "Name EN",
        "Popis CZ",
        "Popis EN",
        "Kategorie",
        "Barva",
        "Materiál",
        "Objem",
        "Cena bez DPH (CZK)",
        "Skladem (ks)",
        "Dodací lhůta (dny)",
        "Značka",
        "Země původu",
        "Štítky",
    ]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F6B3A")
    for it in items:
        code = "" if it.get("_broken") == "code" else it["supplier_sku"]
        cap = f"{it['capacity_ml'] / 1000:.2f} l".replace(".", ",") if it["capacity_ml"] else ""
        ws.append(
            [
                code,
                it["names"]["cs"],
                it["names"]["sk"],
                it["names"]["en"],
                it["descs"]["cs"],
                it["descs"]["en"] if rng.random() < 0.7 else "",
                cats[it["subcategory"]],
                COLOR_NAMES[it["color"]][1],
                MATERIAL_NAMES[it["material"]][1],
                cap,
                round(it["price_eur"] * CZK_PER_EUR, 0),
                it["stock"],
                it["lead"],
                it["brand"],
                it["origin"],
                ", ".join(it["tags"]),
            ]
        )
    ws2 = wb.create_sheet("Info")
    ws2.append(["Dodavatel", "EcoLine Supplies s.r.o."])
    ws2.append(["Měna", "CZK"])
    ws2.append(["Poznámka", "Ceny jsou uvedeny bez DPH. Minimální odběr 50 ks."])
    wb.save(SUPPLIERS_DIR / "supplier_3.xlsx")


def main() -> None:
    SUPPLIERS_DIR.mkdir(parents=True, exist_ok=True)
    PRODUCTS_DIR.mkdir(parents=True, exist_ok=True)
    items = build_catalogue()
    by_supplier = {s: [i for i in items if i["supplier"] == s] for s in (1, 2, 3)}

    # A few deliberately broken rows so the ingestion validator has something to reject.
    for it, kind in zip(rng.sample(by_supplier[1], 3), ["price", "colour", "price"], strict=True):
        it["_broken"] = kind
    for it in rng.sample(by_supplier[2], 2):
        it["_broken"] = "price"
    for it in rng.sample(by_supplier[3], 2):
        it["_broken"] = "code"

    write_supplier_1(by_supplier[1])
    write_supplier_2(by_supplier[2])
    write_supplier_3(by_supplier[3])

    # ERP crosswalk: existing mapping for ~96 % of valid items; the rest are "new" products.
    valid = [i for i in items if not i.get("_broken")]
    order = valid[:]
    rng.shuffle(order)
    hero = next(i for i in order if i["supplier_sku"] == "29904.01")
    order.remove(hero)
    order.insert(41, hero)  # -> NP-10042
    supplier_codes = {1: "EGT", 2: "PTX", 3: "ECO"}
    crosswalk, stock_rows = [], []
    for idx, it in enumerate(order):
        if idx != 41 and rng.random() < 0.04:
            continue  # not yet in the ERP -> ingestion allocates a new SKU
        np_sku = f"NP-{10001 + idx}"
        crosswalk.append(
            {
                "supplier_code": supplier_codes[it["supplier"]],
                "supplier_sku": it["supplier_sku"],
                "sku": np_sku,
            }
        )
        ba = rng.choice([0, 0, 15, 60, 150, 400])
        za = rng.choice([0, 0, 0, 20, 80])
        if np_sku == "NP-10042":
            ba, za = 340, 120
        stock_rows.append(
            {"sku": np_sku, "BA": ba, "ZA": za, "reserved": min(ba, rng.choice([0, 0, 5, 20, 40]))}
        )
    with open(PRODUCTS_DIR / "sku_crosswalk.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["supplier_code", "supplier_sku", "sku"])
        w.writeheader()
        w.writerows(crosswalk)
    with open(PRODUCTS_DIR / "warehouse_stock.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["sku", "BA", "ZA", "reserved"])
        w.writeheader()
        w.writerows(stock_rows)

    # Run the real normalisation pipeline and export the canonical catalogue for reference.
    from services.ingestion.suppliers import load_all_feeds

    report = load_all_feeds(ROOT / "data" / "suppliers")
    fields = [
        "supplier_code",
        "supplier_sku",
        "name_sk",
        "name_en",
        "name_cs",
        "category",
        "subcategory",
        "color",
        "material",
        "capacity_ml",
        "sizes",
        "price_eur",
        "stock",
        "brand",
        "country_of_origin",
        "lead_time_days",
        "tags",
        "description_sk",
        "description_en",
        "description_cs",
    ]
    with open(PRODUCTS_DIR / "products_normalized.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        for p in report.products:
            row = p.model_dump(include=set(fields))
            row["sizes"] = "|".join(p.sizes)
            row["tags"] = "|".join(p.tags)
            w.writerow(row)
    print(
        f"Generated {len(items)} supplier items ({len(by_supplier[1])} XML, {len(by_supplier[2])} JSON, "
        f"{len(by_supplier[3])} XLSX); normalised {len(report.products)}, rejected {len(report.rejected)}; "
        f"crosswalk entries {len(crosswalk)}"
    )


if __name__ == "__main__":
    main()
