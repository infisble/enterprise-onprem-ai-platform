"""Supplier feed ingestion: schema-specific adapters -> one canonical `NormalizedProduct`.

Every supplier ships a different schema (XML with German colour names and comma decimals, JSON
with prices in cents and per-size stock, XLSX with Czech headers and CZK prices). Adapters keep
this mess at the edge of the system; everything downstream (SQL, search, RAG, agent) only sees
the canonical model. Invalid rows are rejected with a reason instead of silently imported.
"""

from __future__ import annotations

import json
import re
import xml.etree.ElementTree as ET
from collections.abc import Iterator
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from pydantic import ValidationError

from packages.core.text import COLOR_LOOKUP, MATERIAL_NAMES, normalize
from packages.schemas.products import NormalizedProduct

CZK_PER_EUR = 25.0  # in production: daily ECB rate from the ERP


@dataclass
class RawRecord:
    supplier_code: str
    source_ref: str  # e.g. "supplier_3.xlsx!Produkty:row 124" – used in rejection reports
    data: dict[str, Any]


@dataclass
class IngestionReport:
    products: list[NormalizedProduct] = field(default_factory=list)
    rejected: list[dict[str, str]] = field(default_factory=list)
    per_supplier: dict[str, dict[str, int]] = field(default_factory=dict)

    def summary(self) -> dict[str, Any]:
        return {
            "normalized": len(self.products),
            "rejected": len(self.rejected),
            "per_supplier": self.per_supplier,
            "rejections": self.rejected[:20],
        }


# ------------------------------------------------------------------------------------------------
def parse_decimal(value: Any) -> float:
    if isinstance(value, int | float):
        return float(value)
    text = str(value).strip().replace(" ", "").replace(" ", "")
    if not re.fullmatch(r"-?\d+([.,]\d+)?", text):
        raise ValueError(f"not a number: {value!r}")
    return float(text.replace(",", "."))


def parse_volume_ml(value: Any) -> int | None:
    if value in (None, ""):
        return None
    m = re.match(r"\s*(\d+(?:[.,]\d+)?)\s*(ml|l)\b", str(value).lower())
    if not m:
        raise ValueError(f"unparseable volume: {value!r}")
    number = float(m.group(1).replace(",", "."))
    return int(round(number * 1000 if m.group(2) == "l" else number))


def canonical_color(value: str | None) -> str:
    if not value:
        raise ValueError("missing colour")
    for token in normalize(value).replace("-", " ").split():
        if token in COLOR_LOOKUP:
            return COLOR_LOOKUP[token]
    raise ValueError(f"unknown colour: {value!r}")


_CS_MATERIALS = {normalize(cs): en for en, (_, cs) in MATERIAL_NAMES.items()}


class SupplierAdapter:
    code: str = ""
    name: str = ""
    country: str = ""
    feed_format: str = ""
    filename: str = ""

    def parse(self, path: Path) -> Iterator[RawRecord]:
        raise NotImplementedError

    def to_canonical(self, raw: RawRecord) -> dict[str, Any]:
        raise NotImplementedError

    def normalize(self, raw: RawRecord) -> NormalizedProduct:
        return NormalizedProduct(**self.to_canonical(raw))


class EuroGiftXmlAdapter(SupplierAdapter):
    code, name, country, feed_format, filename = "EGT", "EuroGift Trading GmbH", "DE", "xml", "supplier_1.xml"
    CATEGORY_MAP = {
        "Drinkware > Vacuum Flasks & Bottles": ("drinkware", "bottles"),
        "Drinkware > Mugs": ("drinkware", "mugs"),
        "Drinkware > Travel Mugs": ("drinkware", "tumblers"),
        "Rain > Umbrellas": ("umbrellas", "umbrellas"),
        "Bags > Rucksacks": ("bags", "backpacks"),
        "Bags > Document Bags": ("bags", "conference bags"),
        "Writing > Ballpens": ("writing", "pens"),
        "Electronics > Powerbanks": ("technology", "power banks"),
        "Electronics > Chargers": ("technology", "chargers"),
        "Electronics > Audio": ("technology", "speakers"),
    }

    def parse(self, path: Path) -> Iterator[RawRecord]:
        root = ET.parse(path).getroot()
        for idx, item in enumerate(root.findall("item"), start=1):
            yield RawRecord(self.code, f"{path.name}:item[{idx}] code={item.get('code')}", {"el": item})

    def to_canonical(self, raw: RawRecord) -> dict[str, Any]:
        el: ET.Element = raw.data["el"]
        names = {n.get("lang"): (n.text or "") for n in el.findall("names/name")}
        descs = {d.get("lang"): (d.text or "") for d in el.findall("descriptions/description")}
        attrs = {a.get("name"): (a.text or "") for a in el.findall("attributes/attr")}
        prices = {int(p.get("qty", "1")): p.text for p in el.findall("pricing/price")}
        category, sub = self.CATEGORY_MAP[el.find("category").get("path", "")]  # type: ignore[union-attr]
        return {
            "supplier_code": self.code,
            "supplier_sku": el.get("code"),
            "name_en": names.get("en", ""),
            "name_sk": names.get("sk", ""),
            "description_en": descs.get("en", ""),
            "description_sk": descs.get("sk", ""),
            "category": category,
            "subcategory": sub,
            "color": canonical_color(attrs.get("colour")),
            "material": attrs.get("material", "").strip().lower(),
            "capacity_ml": parse_volume_ml(attrs.get("volume")),
            "price_eur": parse_decimal(prices.get(1)),
            "stock": int(el.findtext("stock", "0")),
            "brand": el.findtext("brand", ""),
            "country_of_origin": el.findtext("origin", ""),
            "lead_time_days": int(el.findtext("leadtime", "10")),
            "tags": [t for t in el.findtext("keywords", "").split(";") if t],
        }

    @staticmethod
    def price_tiers(el: ET.Element) -> list[tuple[int, float]]:
        return [(int(p.get("qty", "1")), parse_decimal(p.text)) for p in el.findall("pricing/price")]


class PromoTexJsonAdapter(SupplierAdapter):
    code, name, country, feed_format, filename = "PTX", "PromoTex s.r.o.", "CZ", "json", "supplier_2.json"
    GROUP_MAP = {
        "T-Shirts": ("apparel", "t-shirts"),
        "Polos": ("apparel", "polo shirts"),
        "Sweats & Hoodies": ("apparel", "hoodies"),
        "Headwear": ("apparel", "caps"),
        "Bags": ("bags", "tote bags"),
        "Accessories": ("lifestyle", "lanyards"),
    }

    def parse(self, path: Path) -> Iterator[RawRecord]:
        feed = json.loads(path.read_text(encoding="utf-8"))
        codes = feed.get("colour_codes", {})
        for idx, p in enumerate(feed["products"]):
            yield RawRecord(
                self.code, f"{path.name}:$.products[{idx}] id={p.get('id')}", {"p": p, "codes": codes}
            )

    @staticmethod
    def _material(fabric: str) -> str:
        f = fabric.lower()
        if "recycled" in f or "rpet" in f:
            return "recycled pet"
        if "organic" in f:
            return "organic cotton"
        if "cotton" in f:
            return "cotton"
        if "polyester" in f:
            return "polyester"
        raise ValueError(f"unknown fabric: {fabric!r}")

    def to_canonical(self, raw: RawRecord) -> dict[str, Any]:
        p, codes = raw.data["p"], raw.data["codes"]
        category, sub = self.GROUP_MAP[p["group"]]
        sizes = [s["size"] for s in p.get("sizes", []) if s["size"] != "ONE"]
        cents = int(p["price"]["amount_cents"])
        if cents <= 0:
            raise ValueError(f"non-positive price: {cents} cents")
        return {
            "supplier_code": self.code,
            "supplier_sku": p["id"],
            "name_en": p["title"].get("en", ""),
            "name_sk": p["title"].get("sk", ""),
            "name_cs": p["title"].get("cs", ""),
            "description_en": p.get("body", {}).get("en", ""),
            "description_cs": p.get("body", {}).get("cs", ""),
            "category": category,
            "subcategory": sub,
            "color": canonical_color(codes.get(p["colour_code"], p["colour_code"])),
            "material": self._material(p["fabric"]),
            "sizes": sizes,
            "price_eur": round(cents / 100, 2),
            "stock": sum(int(s.get("stock", 0)) for s in p.get("sizes", [])),
            "brand": p.get("brand", ""),
            "country_of_origin": p.get("made_in", ""),
            "lead_time_days": int(p.get("delivery_days", 10)),
            "tags": list(p.get("labels", [])),
        }


class EcoLineXlsxAdapter(SupplierAdapter):
    code, name, country, feed_format, filename = (
        "ECO",
        "EcoLine Supplies s.r.o.",
        "CZ",
        "xlsx",
        "supplier_3.xlsx",
    )
    CATEGORY_MAP = {
        "Kancelář / Zápisníky": ("office", "notebooks"),
        "Nápoje / Hrnky": ("drinkware", "mugs"),
        "Nápoje / Láhve": ("drinkware", "bottles"),
        "Psací potřeby / Tužky": ("writing", "pencils"),
        "Tašky / Nákupní tašky": ("bags", "tote bags"),
        "Elektronika / Nabíječky": ("technology", "chargers"),
        "Elektronika / Powerbanky": ("technology", "power banks"),
        "Domov / Doplňky": ("lifestyle", "home"),
    }
    # Configurable column mapping: header text -> canonical field (changes per supplier contract).
    COLUMNS = {
        "Kód zboží": "sku",
        "Název CZ": "name_cs",
        "Název SK": "name_sk",
        "Name EN": "name_en",
        "Popis CZ": "description_cs",
        "Popis EN": "description_en",
        "Kategorie": "category",
        "Barva": "color",
        "Materiál": "material",
        "Objem": "volume",
        "Cena bez DPH (CZK)": "price_czk",
        "Skladem (ks)": "stock",
        "Dodací lhůta (dny)": "lead",
        "Značka": "brand",
        "Země původu": "origin",
        "Štítky": "tags",
    }

    def parse(self, path: Path) -> Iterator[RawRecord]:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb["Produkty"]
        rows = ws.iter_rows(values_only=True)
        headers = [str(h) for h in next(rows)]
        for idx, row in enumerate(rows, start=2):
            data = {self.COLUMNS[h]: v for h, v in zip(headers, row, strict=False) if h in self.COLUMNS}
            yield RawRecord(self.code, f"{path.name}!Produkty:row {idx}", data)
        wb.close()

    def to_canonical(self, raw: RawRecord) -> dict[str, Any]:
        d = raw.data
        if not d.get("sku"):
            raise ValueError("missing supplier product code")
        category, sub = self.CATEGORY_MAP[d["category"]]
        material = _CS_MATERIALS.get(normalize(str(d["material"])))
        if not material:
            raise ValueError(f"unknown material: {d['material']!r}")
        return {
            "supplier_code": self.code,
            "supplier_sku": str(d["sku"]),
            "name_cs": d.get("name_cs") or "",
            "name_sk": d.get("name_sk") or "",
            "name_en": d.get("name_en") or "",
            "description_cs": d.get("description_cs") or "",
            "description_en": d.get("description_en") or "",
            "category": category,
            "subcategory": sub,
            "color": canonical_color(str(d["color"])),
            "material": material,
            "capacity_ml": parse_volume_ml(d.get("volume")),
            "price_eur": round(parse_decimal(d["price_czk"]) / CZK_PER_EUR, 2),
            "stock": int(d.get("stock") or 0),
            "brand": d.get("brand") or "",
            "country_of_origin": d.get("origin") or "",
            "lead_time_days": int(d.get("lead") or 10),
            "tags": [t.strip() for t in str(d.get("tags") or "").split(",") if t.strip()],
        }


ADAPTERS: list[SupplierAdapter] = [EuroGiftXmlAdapter(), PromoTexJsonAdapter(), EcoLineXlsxAdapter()]


def load_feed(adapter: SupplierAdapter, path: Path, report: IngestionReport) -> None:
    ok = bad = 0
    for raw in adapter.parse(path):
        try:
            report.products.append(adapter.normalize(raw))
            ok += 1
        except (ValueError, KeyError, ValidationError, AttributeError) as exc:
            reason = (
                str(exc).splitlines()[0]
                if not isinstance(exc, ValidationError)
                else "; ".join(f"{'.'.join(map(str, e['loc']))}: {e['msg']}" for e in exc.errors())
            )
            report.rejected.append({"supplier": adapter.code, "source": raw.source_ref, "reason": reason})
            bad += 1
    report.per_supplier[adapter.code] = {"normalized": ok, "rejected": bad}


def load_all_feeds(suppliers_dir: Path) -> IngestionReport:
    report = IngestionReport()
    for adapter in ADAPTERS:
        load_feed(adapter, suppliers_dir / adapter.filename, report)
    return report
