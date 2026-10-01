"""File-type-aware parsing. Each parser returns `Section`s that keep their *locator*
(page / heading / sheet+row / JSON path / XML path) so every chunk can be cited precisely.

| Type      | Unit of meaning                   | Locator                     |
|-----------|-----------------------------------|-----------------------------|
| PDF       | heading-delimited sections        | page + section              |
| DOCX      | Heading-styled sections           | section                     |
| XLSX      | one row -> one record (template)  | sheet + row                 |
| XML       | repeated element -> one record    | xpath-like path             |
| JSON      | array item -> one record          | JSONPath                    |
| MD / TXT  | markdown heading sections         | section                     |
"""

from __future__ import annotations

import json
import re
import xml.etree.ElementTree as ET
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from docx import Document as Docx
from openpyxl import load_workbook
from pypdf import PdfReader


@dataclass
class Section:
    text: str
    locator: dict[str, Any]
    heading: str = ""
    kind: str = "prose"  # prose | record
    meta: dict[str, Any] = field(default_factory=dict)


_NUM_HEADING = re.compile(r"^(\d+(\.\d+)*)\.?\s+\S.{0,80}$")


def _is_heading(line: str) -> bool:
    line = line.strip()
    if not line or len(line) > 90 or line.endswith((".", ",", ":", ";")):
        return False
    if _NUM_HEADING.match(line):
        return True
    words = line.split()
    return (
        1 <= len(words) <= 8
        and line[0].isupper()
        and sum(w[0].isupper() for w in words if w[0].isalpha()) >= 1
        and not any(ch.isdigit() for ch in line[-3:])
        and len(line) < 60
    )


def parse_pdf(path: Path) -> list[Section]:
    reader = PdfReader(str(path))
    sections: list[Section] = []
    heading, buf, start_page = "", list[str](), 1
    for page_no, page in enumerate(reader.pages, start=1):
        lines = (page.extract_text() or "").splitlines()
        for idx, line in enumerate(lines):
            is_title = page_no == 1 and idx == 0
            if is_title or _is_heading(line):
                if buf:
                    sections.append(Section(" ".join(buf), {"page": start_page, "section": heading}, heading))
                heading, buf, start_page = line.strip(), [], page_no
            else:
                if not buf:
                    start_page = page_no
                buf.append(line.strip())
    if buf:
        sections.append(Section(" ".join(buf), {"page": start_page, "section": heading}, heading))
    return sections


def parse_docx(path: Path) -> list[Section]:
    doc = Docx(str(path))
    sections: list[Section] = []
    heading, buf = "", list[str]()
    for para in doc.paragraphs:
        style = (para.style.name if para.style is not None else "") or ""
        text = para.text.strip()
        if not text:
            continue
        if style.startswith(("Heading", "Title")):
            if buf:
                sections.append(Section("\n".join(buf), {"section": heading}, heading))
            heading, buf = text, []
        else:
            buf.append(text)
    if buf:
        sections.append(Section("\n".join(buf), {"section": heading}, heading))
    return sections


def parse_xlsx(path: Path, sheet_templates: dict[str, str] | None = None) -> list[Section]:
    """Row-to-document transformation. A template like "Category {Category}: margin {Target margin %} %"
    turns a row into a readable sentence; without a template the row becomes "Header: value; ..."."""
    wb = load_workbook(path, read_only=True, data_only=True)
    sections: list[Section] = []
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        try:
            headers = [str(h) if h is not None else f"col{i}" for i, h in enumerate(next(rows))]
        except StopIteration:
            continue
        template = (sheet_templates or {}).get(ws.title)
        for row_no, row in enumerate(rows, start=2):
            record = {h: ("" if v is None else v) for h, v in zip(headers, row, strict=False)}
            if not any(str(v).strip() for v in record.values()):
                continue
            try:
                text = (
                    template.format(**record)
                    if template
                    else "; ".join(f"{k}: {v}" for k, v in record.items() if v != "")
                )
            except (KeyError, IndexError):
                text = "; ".join(f"{k}: {v}" for k, v in record.items() if v != "")
            sections.append(
                Section(text, {"sheet": ws.title, "row": row_no}, ws.title, kind="record", meta=record)
            )
    wb.close()
    return sections


def _flatten(prefix: str, value: Any, out: list[str]) -> None:
    if isinstance(value, dict):
        for k, v in value.items():
            _flatten(f"{prefix}.{k}" if prefix else k, v, out)
    elif isinstance(value, list):
        if all(not isinstance(v, dict | list) for v in value):
            out.append(f"{prefix}: {', '.join(map(str, value))}")
        else:
            for i, v in enumerate(value):
                _flatten(f"{prefix}[{i}]", v, out)
    elif value not in (None, ""):
        out.append(f"{prefix}: {value}")


def parse_json(path: Path) -> list[Section]:
    """Schema-aware: find the main array of objects (e.g. $.products) and emit one record per item."""
    data = json.loads(path.read_text(encoding="utf-8"))
    array_key, items = None, None
    if isinstance(data, list):
        array_key, items = "$", data
    elif isinstance(data, dict):
        for k, v in data.items():
            if isinstance(v, list) and v and all(isinstance(x, dict) for x in v):
                array_key, items = f"$.{k}", v
                break
    if items is None:
        lines: list[str] = []
        _flatten("", data, lines)
        return [Section("\n".join(lines), {"path": "$"}, path.name)]
    sections = []
    for i, item in enumerate(items):
        lines = []
        _flatten("", item, lines)
        sections.append(
            Section("\n".join(lines), {"path": f"{array_key}[{i}]"}, str(item.get("id", i)), kind="record")
        )
    return sections


def parse_xml(path: Path) -> list[Section]:
    """Schema-aware: the most frequent child element of the root is treated as the record type."""
    root = ET.parse(path).getroot()
    tag_counts = Counter(child.tag for child in root)
    if not tag_counts:
        return [Section(" ".join(root.itertext()).strip(), {"path": f"/{root.tag}"}, root.tag)]
    record_tag, _ = tag_counts.most_common(1)[0]
    sections = []
    for i, el in enumerate(root.findall(record_tag), start=1):
        lines = [f"{k}: {v}" for k, v in el.attrib.items()]
        for sub in el.iter():
            if sub is el:
                continue
            label = sub.tag + "".join(f"[{k}={v}]" for k, v in sub.attrib.items())
            if sub.text and sub.text.strip():
                lines.append(f"{label}: {sub.text.strip()}")
            elif sub.attrib and not list(sub):
                lines.append(label)
        sections.append(
            Section(
                "\n".join(lines),
                {"path": f"/{root.tag}/{record_tag}[{i}]"},
                el.get("code", el.get("id", str(i))),
                kind="record",
            )
        )
    return sections


def parse_markdown(path: Path) -> list[Section]:
    sections: list[Section] = []
    heading, buf = "", list[str]()
    for line in path.read_text(encoding="utf-8").splitlines():
        m = re.match(r"^(#{1,6})\s+(.*)", line)
        if m:
            if buf:
                sections.append(Section("\n".join(buf).strip(), {"section": heading}, heading))
            heading, buf = m.group(2).strip(), []
        elif line.strip():
            buf.append(line.strip())
    if buf:
        sections.append(Section("\n".join(buf).strip(), {"section": heading}, heading))
    return sections


PARSERS = {
    ".pdf": parse_pdf,
    ".docx": parse_docx,
    ".xlsx": parse_xlsx,
    ".json": parse_json,
    ".xml": parse_xml,
    ".md": parse_markdown,
    ".txt": parse_markdown,
}


def parse_file(path: Path, options: dict[str, Any] | None = None) -> list[Section]:
    suffix = path.suffix.lower()
    if suffix not in PARSERS:
        raise ValueError(f"unsupported file type: {suffix}")
    if suffix == ".xlsx":
        return parse_xlsx(path, (options or {}).get("sheets"))
    return PARSERS[suffix](path)
